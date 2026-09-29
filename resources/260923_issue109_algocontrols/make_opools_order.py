#!/usr/bin/env python3
"""make_opools_order.py -- design run outputs -> IDT oPools order .xlsx.

Issue algocontrols (IsPETase / LCC positive controls).  Step 8: put the ordered
oligos into IDT's oPools upload format, validated against opoolsentrysample.xlsx.

    order/constructs_simple.csv (per cluster)  ->  opools_order.xlsx

THE FORMAT (from opoolsentrysample.xlsx)
----------------------------------------
Sheet1, exactly two data columns, one row per oligo:
    A  "Pool name"   distinct name per pool; multiple pools allowed in one file
    B  "Sequence"    ACGT + IUPAC mixed bases (Table 1); <= 350 nt; 5'Phos optional
The sample's column-D instruction text is IDT template guidance, not order data,
so it is intentionally omitted from the upload file.

POOLING
-------
One pool per cluster -- all of a cluster's fragment-position oligos go into one
tube for the combinatorial Golden Gate assembly.  Pool names are the cluster
labels (IsPETase, LCC).  Both pools live in one workbook.
"""
import argparse
import csv
import glob
import os
import sys

import openpyxl
from openpyxl.styles import Font

MAX_NT = 350
IUPAC = set("ACGTRYSWKMBDHVN")
FONT = "Arial"

# pool label -> design output folder (newest run is used)
CLUSTERS = [("IsPETase", "ispetase_design"), ("LCC", "lcc_design")]


def newest_constructs(base, design_dir):
    hits = glob.glob(os.path.join(base, design_dir, "*", "order", "constructs_simple.csv"))
    if not hits:
        sys.exit(f"no constructs_simple.csv under {design_dir} -- run make_order.py first")
    return sorted(hits)[-1]


def read_seqs(path):
    with open(path, newline="") as fh:
        rows = list(csv.DictReader(fh))
    out = []
    for r in rows:
        seq = (r.get("sequence_to_order") or list(r.values())[-1]).strip().upper()
        out.append(seq)
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--base", default=os.path.dirname(os.path.abspath(__file__)))
    ap.add_argument("--out", default=None)
    args = ap.parse_args()
    out = args.out or os.path.join(args.base, "opools_order.xlsx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "Pool name"
    ws["B1"] = "Sequence"
    ws["A1"].font = ws["B1"].font = Font(name=FONT, bold=True)

    row = 2
    problems, summary = [], []
    for pool, design_dir in CLUSTERS:
        src = newest_constructs(args.base, design_dir)
        seqs = read_seqs(src)
        for seq in seqs:
            bad = set(seq) - IUPAC
            if bad:
                problems.append(f"{pool} row {row}: non-IUPAC chars {sorted(bad)}")
            if len(seq) > MAX_NT:
                problems.append(f"{pool} row {row}: {len(seq)} nt > {MAX_NT}")
            ws.cell(row=row, column=1, value=pool).font = Font(name=FONT)
            ws.cell(row=row, column=2, value=seq).font = Font(name=FONT)
            row += 1
        summary.append((pool, len(seqs), os.path.relpath(src, args.base)))

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 120

    if problems:
        sys.exit("ABORT -- oPools constraint violations:\n  " + "\n  ".join(problems))

    wb.save(out)
    print(f"wrote {out}")
    total = 0
    for pool, n, src in summary:
        print(f"  pool {pool:<10} {n:>3} oligos   (from {src})")
        total += n
    print(f"  {len(summary)} pools, {total} oligos total")


if __name__ == "__main__":
    main()
