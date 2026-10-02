#!/usr/bin/env python3
"""opools_from_run.py -- finished design run(s) -> one IDT oPools order .xlsx.

Pipeline generalisation of the control-run make_opools_order.py: instead of a
hard-coded two-cluster list, it takes any number of pool=run_dir pairs and writes
them as pools in one workbook (one pool per cluster), validated against the oPools
constraints.

    python3 opools_from_run.py --out order.xlsx <pool>=<run_dir> [<pool>=<run_dir> ...]

Each <run_dir> must contain order/constructs_simple.csv (make_order.py output).
Format (from opoolsentrysample.xlsx): Sheet1, columns "Pool name" | "Sequence",
one row per oligo; ACGT + IUPAC mixed bases; <= 350 nt.
"""
import argparse
import csv
import os
import sys

import openpyxl
from openpyxl.styles import Font

MAX_NT = 350
IUPAC = set("ACGTRYSWKMBDHVN")
FONT = "Arial"


def read_seqs(run_dir):
    path = os.path.join(run_dir, "order", "constructs_simple.csv")
    if not os.path.exists(path):
        sys.exit(f"missing {path} -- run make_order.py on {run_dir} first")
    with open(path, newline="") as fh:
        rows = list(csv.DictReader(fh))
    return [(r.get("sequence_to_order") or list(r.values())[-1]).strip().upper() for r in rows]


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", required=True)
    ap.add_argument("pairs", nargs="+", help="pool=run_dir (one per cluster)")
    args = ap.parse_args()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"], ws["B1"] = "Pool name", "Sequence"
    ws["A1"].font = ws["B1"].font = Font(name=FONT, bold=True)

    row, problems, summary = 2, [], []
    for pair in args.pairs:
        if "=" not in pair:
            sys.exit(f"bad pair {pair!r} -- expected pool=run_dir")
        pool, run_dir = pair.split("=", 1)
        seqs = read_seqs(run_dir)
        for seq in seqs:
            bad = set(seq) - IUPAC
            if bad:
                problems.append(f"{pool} row {row}: non-IUPAC {sorted(bad)}")
            if len(seq) > MAX_NT:
                problems.append(f"{pool} row {row}: {len(seq)} nt > {MAX_NT}")
            ws.cell(row=row, column=1, value=pool).font = Font(name=FONT)
            ws.cell(row=row, column=2, value=seq).font = Font(name=FONT)
            row += 1
        summary.append((pool, len(seqs)))

    if problems:
        sys.exit("ABORT -- oPools constraint violations:\n  " + "\n  ".join(problems))

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 120
    wb.save(args.out)

    total = sum(n for _, n in summary)
    print(f"wrote {args.out}")
    for pool, n in summary:
        print(f"  pool {pool:<12} {n:>4} oligos")
    print(f"  {len(summary)} pools, {total} oligos total")


if __name__ == "__main__":
    main()
